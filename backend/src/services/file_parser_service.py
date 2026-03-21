# -*- coding: utf-8 -*-
"""
文件解析服务

提供各种文件格式的文本提取功能
"""
import logging
import os
from typing import Set
from pathlib import Path

logger = logging.getLogger(__name__)


class FileParserService:
    """文件解析服务 - 提取文件文本内容"""

    SUPPORTED_FORMATS: Set[str] = {
        '.txt', '.md', '.json',      # 纯文本
        '.pdf',                       # PDF
        '.doc', '.docx',             # Word
        '.xls', '.xlsx',             # Excel
        '.ppt', '.pptx',             # PPT
    }

    async def parse_file(self, file_path: str, filename: str) -> str:
        """
        解析文件，提取文本内容

        Args:
            file_path: 文件路径
            filename: 文件名

        Returns:
            str: 提取的文本内容

        Raises:
            ValueError: 不支持的文件格式
        """
        ext = Path(filename).suffix.lower()

        if ext in {'.txt', '.md', '.json'}:
            return await self._parse_text(file_path)
        elif ext == '.pdf':
            return await self._parse_pdf(file_path)
        elif ext in {'.doc', '.docx'}:
            return await self._parse_word(file_path)
        elif ext in {'.xls', '.xlsx'}:
            return await self._parse_excel(file_path)
        elif ext in {'.ppt', '.pptx'}:
            return await self._parse_ppt(file_path)
        else:
            raise ValueError(f"Unsupported file format: {ext}")

    async def _parse_text(self, file_path: str) -> str:
        """解析纯文本文件"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except UnicodeDecodeError:
            # 尝试其他编码
            with open(file_path, 'r', encoding='gb18030') as f:
                return f.read()

    async def _parse_pdf(self, file_path: str) -> str:
        """解析 PDF 文件"""
        try:
            import PyPDF2
        except ImportError:
            logger.error("PyPDF2 未安装，请运行: pip install PyPDF2")
            return "[PDF 解析失败: PyPDF2 未安装]"

        try:
            text_parts = []
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    text = page.extract_text()
                    if text:
                        text_parts.append(text)
            return '\n'.join(text_parts)
        except Exception as e:
            logger.error(f"PDF 解析失败: {e}")
            return f"[PDF 解析失败: {str(e)}]"

    async def _parse_word(self, file_path: str) -> str:
        """解析 Word 文件"""
        try:
            from docx import Document
        except ImportError:
            logger.error("python-docx 未安装，请运行: pip install python-docx")
            return "[Word 解析失败: python-docx 未安装]"

        try:
            doc = Document(file_path)
            text_parts = []
            for paragraph in doc.paragraphs:
                if paragraph.text:
                    text_parts.append(paragraph.text)
            return '\n'.join(text_parts)
        except Exception as e:
            logger.error(f"Word 解析失败: {e}")
            return f"[Word 解析失败: {str(e)}]"

    async def _parse_excel(self, file_path: str) -> str:
        """解析 Excel 文件"""
        try:
            import openpyxl
        except ImportError:
            logger.error("openpyxl 未安装，请运行: pip install openpyxl")
            return "[Excel 解析失败: openpyxl 未安装]"

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            text_parts = []
            for sheet in wb.worksheets:
                text_parts.append(f"### 工作表: {sheet.title} ###")
                for row in sheet.iter_rows(values_only=True):
                    row_text = '\t'.join(str(cell) if cell is not None else '' for cell in row)
                    if row_text.strip():
                        text_parts.append(row_text)
            return '\n'.join(text_parts)
        except Exception as e:
            logger.error(f"Excel 解析失败: {e}")
            return f"[Excel 解析失败: {str(e)}]"

    async def _parse_ppt(self, file_path: str) -> str:
        """解析 PPT 文件"""
        try:
            from pptx import Presentation
        except ImportError:
            logger.error("python-pptx 未安装，请运行: pip install python-pptx")
            return "[PPT 解析失败: python-pptx 未安装]"

        try:
            prs = Presentation(file_path)
            text_parts = []
            for i, slide in enumerate(prs.slides):
                text_parts.append(f"### 幻灯片 {i + 1} ###")
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        text_parts.append(shape.text)
            return '\n'.join(text_parts)
        except Exception as e:
            logger.error(f"PPT 解析失败: {e}")
            return f"[PPT 解析失败: {str(e)}]"
